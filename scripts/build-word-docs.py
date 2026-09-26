"""Builds docs/FreshFind_Project_Report.docx and ReadMe.doc from the Markdown files in this project.

Run it from anywhere with:  python scripts/build-word-docs.py
It needs Python 3 with python-docx and Pillow:  pip install python-docx pillow
With Microsoft Word installed it also fills in the report's table of contents and saves the
ReadMe in the Word 97-2003 format. Without Word it tries LibreOffice, and otherwise keeps ReadMe.docx.
"""

import io
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
TEAM_NOTES = ROOT / "team-notes"

REPORT_PATH = DOCS / "FreshFind_Project_Report.docx"
README_DOCX = ROOT / "ReadMe.docx"
README_DOC = ROOT / "ReadMe.doc"

LIVE_SITE = "https://geraldxtra.github.io/freshfind/"
TEAM_NAME = "Tech Hive"

FONT = "Calibri"
GREEN = RGBColor(0x1A, 0x56, 0x32)
GREEN_LIGHT = RGBColor(0x1F, 0x6B, 0x3D)
TEXT = RGBColor(0x16, 0x21, 0x1B)
MUTED = RGBColor(0x5C, 0x6B, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINK_HEX = "0563C1"
HEADER_FILL = "1A5632"
BORDER_HEX = "C9D3CC"
SCORE_FILLS = [(90, "E3F2E6"), (50, "FDF0DA"), (0, "FBE1DE")]

PAGE_SHORT = Cm(21.0)
PAGE_LONG = Cm(29.7)
MARGIN = Cm(2.5)
PORTRAIT_WIDTH = 16.0
LANDSCAPE_WIDTH = 24.7
PORTRAIT_IMAGE_HEIGHT = 20.5
LANDSCAPE_IMAGE_HEIGHT = 12.5
DIAGRAM_DPI = 400
PHOTO_DPI = 300

ROLE_TEXT = {
    "Eberechukwu Uchechukwu Gerald, Team Leader": (
        "Responsible for the Home page, the chatbot and its rule engine, the Bookmarks page and the Not Found page. "
        "Also responsible for the data files and the shared frame: the theme, Navbar, Footer, search overlay, market card, "
        "bookmark button, breadcrumbs, helpers and the bookmarks context. Planned the project from the SRS, chose the stack "
        "and the design direction, set up the repository, the branches and the working rules, and wrote the team "
        "documentation. Reviews every pull request and runs the integration testing, Lighthouse, the report and the demo video."
    ),
    "Chukwujekwu Chimdiuso Amanda": (
        "Responsible for the Market Directory and the Contact Us page. The Market Directory is the biggest list page on the "
        "site: filtering by area, day, produce type and open now, three sort orders including distance after location is "
        "allowed, filters arriving from the address bar, the results count and the empty state. The Contact page has the "
        "live location map that handles allowed, blocked and unsupported states."
    ),
    "Ibrahim Ogunsola Kelvin": (
        "Responsible for the Market Detail page and the About Us page. Market Detail is the deepest page on the site: one "
        "market loaded from the id in the address, the banner with the live open badge and save button, the weekly schedule "
        "table with today highlighted, the typical produce tiles pulled from the produce data, the map, the next opening time "
        "and the directions link. The About page has the mission, the three reasons and the team cards."
    ),
    "Uyi Osakue Uhunwa": (
        "Responsible for the Produce Guide and the Seasonal Picks page. The Produce Guide has the five category tabs, the "
        "search box, the season badges and the Found at links into market pages. Seasonal Picks has the banner, this week's "
        "top picks computed from the current month, and the season calendar that draws a bar for every month each item is in "
        "season."
    ),
}

REWORDS = {
    "Each member has a walkthrough of their own pages in docs/members, and we read each other's pull requests before merging.":
        "Each member wrote a walkthrough of their own pages, and we read each other's pull requests before merging.",
    "The full report for the Home page after the fixes is saved in the lighthouse folder next to this file, as home-desktop.html.":
        "The full report for the Home page after the fixes is saved in the docs/lighthouse folder of the source code, as home-desktop.html.",
}

PRIVATE_NAMES = [
    "TEAM_GUIDE", "TEAM_ROLES", "VIDEO_SCRIPT", "THEME_GUIDE", "IMAGES_GUIDE",
    "ASSUMPTIONS.md", "DATA_GUIDE", "walkthroughs", "docs/members",
]
CODE_PATTERNS = [r"=>", r"[{}]", r"</?[A-Za-z][^>]*>", r"\bconst\b", r"\bexport (default|function|const)\b",r"\w+\(\)", r"\w+\.\w+\("]

URL_PATTERN = re.compile(r"https://[^\s]+")


def read(path):
    return path.read_text(encoding="utf-8").replace("\r\n", "\n")


def reword(text):
    for old, new in REWORDS.items():
        text = text.replace(old, new)
    return text


def split_sections(text, level):
    marker = "#" * level + " "
    found = {}
    current = None
    lines = []
    for line in text.split("\n"):
        if line.startswith(marker):
            if current is not None:
                found[current] = "\n".join(lines).strip("\n")
            current = line[len(marker):].strip()
            lines = []
        elif current is not None:
            lines.append(line)
    if current is not None:
        found[current] = "\n".join(lines).strip("\n")
    return found


def intro_of(text):
    lines = []
    for line in text.split("\n"):
        if line.startswith("#"):
            if lines:
                break
            continue
        lines.append(line)
    return "\n".join(lines).strip("\n")


def parse_blocks(text):
    blocks = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            blocks.append(("blank", None))
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)$", line)
        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", line)
        number = re.match(r"^(\d+)\.\s+(.*)$", line)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
        elif image:
            blocks.append(("image", (image.group(1), image.group(2))))
        elif line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
        elif line.startswith("- "):
            if blocks and blocks[-1][0] == "bullets":
                blocks[-1][1].append(line[2:].strip())
            else:
                blocks.append(("bullets", [line[2:].strip()]))
        elif number:
            if blocks and blocks[-1][0] == "numbers":
                blocks[-1][1].append((number.group(1), number.group(2).strip()))
            else:
                blocks.append(("numbers", [(number.group(1), number.group(2).strip())]))
        else:
            blocks.append(("paragraph", line))
    return [block for block in blocks if block[0] != "blank"]


def force_font(element, name):
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attribute)) is not None:
            del fonts.attrib[qn(attribute)]
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def style_text(style, size, bold=False, italic=False, color=TEXT):
    force_font(style.element, FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color


def paragraph_style(styles, name):
    if name in [style.name for style in styles]:
        return styles[name]
    return styles.add_style(name, 1)


def setup_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    style_text(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    heading_specs = [
        ("Heading 1", 20, GREEN, Pt(0), Pt(14), True),
        ("Heading 2", 14, GREEN_LIGHT, Pt(16), Pt(6), False),
        ("Heading 3", 12, TEXT, Pt(12), Pt(4), False),
    ]
    for name, size, color, before, after, page_break in heading_specs:
        style = styles[name]
        style_text(style, size, bold=True, color=color)
        style.paragraph_format.space_before = before
        style.paragraph_format.space_after = after
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = page_break
        style.paragraph_format.line_spacing = 1.0

    bullet = styles["List Bullet"]
    style_text(bullet, 11)
    bullet.paragraph_format.space_after = Pt(3)

    caption = paragraph_style(styles, "Figure Caption")
    caption.base_style = normal
    style_text(caption, 9, italic=True, color=MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)

    table_text = paragraph_style(styles, "Table Text")
    table_text.base_style = normal
    style_text(table_text, 10)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0


def set_page(section, landscape):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = PAGE_LONG if landscape else PAGE_SHORT
    section.page_height = PAGE_SHORT if landscape else PAGE_LONG
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = MARGIN
    section.header_distance = section.footer_distance = Cm(1.2)


def new_section(document, landscape):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.different_first_page_header_footer = False
    set_page(section, landscape)
    return section


def add_field(paragraph, instruction, placeholder):
    begin = paragraph.add_run()
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin._r.append(begin_char)
    code = paragraph.add_run()
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = f" {instruction} "
    code._r.append(instruction_text)
    separate = paragraph.add_run()
    separate_char = OxmlElement("w:fldChar")
    separate_char.set(qn("w:fldCharType"), "separate")
    separate._r.append(separate_char)
    paragraph.add_run(placeholder)
    end = paragraph.add_run()
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end._r.append(end_char)


def add_hyperlink(paragraph, url, text, bold=False, size=None):
    relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    if bold:
        properties.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_HEX)
    properties.append(color)
    if size:
        font_size = OxmlElement("w:sz")
        font_size.set(qn("w:val"), str(int(size * 2)))
        properties.append(font_size)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)
    link.append(run)
    paragraph._p.append(link)


def add_text(paragraph, text, bold=False, italic=False, size=None, color=None):
    position = 0
    for match in URL_PATTERN.finditer(text):
        url = match.group(0).rstrip(".,;:)")
        start = match.start()
        if start > position:
            styled_run(paragraph, text[position:start], bold, italic, size, color)
        add_hyperlink(paragraph, url, url, bold=bold, size=size)
        position = start + len(url)
    if position < len(text):
        styled_run(paragraph, text[position:], bold, italic, size, color)


def styled_run(paragraph, text, bold=False, italic=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold or None
    run.italic = italic or None
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def add_paragraph(document, text, style=None, lead=None, align=None):
    paragraph = document.add_paragraph(style=style)
    if lead:
        add_text(paragraph, lead, bold=True)
        text = " " + text if text else ""
    add_text(paragraph, text)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def split_lead(text, separator):
    index = text.find(separator)
    if index == -1:
        return None, text
    cut = index + len(separator.rstrip())
    return text[:cut], text[cut:].strip()


def add_bullets(document, items, lead_separator=None):
    for item in items:
        lead, rest = split_lead(item, lead_separator) if lead_separator else (None, item)
        add_paragraph(document, rest, style="List Bullet", lead=lead)


def add_numbers(document, items, lead_separator=None):
    for number, item in items:
        lead, rest = split_lead(item, lead_separator) if lead_separator else (None, item)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.75)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(f"{number}.\t")
        if lead:
            add_text(paragraph, lead, bold=True)
            rest = " " + rest
        add_text(paragraph, rest)


def cell_fill(cell, hex_color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    properties.append(shading)


def table_borders(table, hex_color):
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        if hex_color:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), hex_color)
        else:
            element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def row_flags(row, header=False):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))
    if header:
        properties.append(OxmlElement("w:tblHeader"))


def score_fill(text):
    if not text.isdigit():
        return None
    value = int(text)
    for floor, color in SCORE_FILLS:
        if value >= floor:
            return color
    return None


def add_table(document, header, rows, widths, size=10, scores=False):
    table = document.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table, BORDER_HEX)
    row_flags(table.rows[0], header=True)
    for index, title in enumerate(header):
        cell = table.rows[0].cells[index]
        cell_fill(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.style = "Table Text"
        styled_run(paragraph, title, bold=True, size=size, color=WHITE)
    for values in rows:
        row = table.add_row()
        row_flags(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table Text"
            add_text(paragraph, value, size=size)
            fill = score_fill(value) if scores else None
            if fill:
                cell_fill(cell, fill)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Cm(width)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


class Figures:
    def __init__(self):
        self.count = 0

    def next(self):
        self.count += 1
        return self.count


def fitted_size(path, max_width_cm, max_height_cm):
    with Image.open(path) as image:
        width, height = image.size
    scale = min(max_width_cm / width, max_height_cm / height)
    return width * scale, height * scale


def image_stream(path, width_cm, dpi, photo):
    with Image.open(path) as image:
        target = math.ceil(width_cm / 2.54 * dpi)
        picture = image.convert("RGB")
        if picture.width > target:
            height = round(picture.height * target / picture.width)
            picture = picture.resize((target, height), Image.LANCZOS)
        stream = io.BytesIO()
        if photo:
            picture.save(stream, format="JPEG", quality=85, optimize=True)
        else:
            picture.save(stream, format="PNG", optimize=True)
    stream.seek(0)
    return stream


def add_picture(container_paragraph, path, width_cm, height_cm, dpi, photo):
    stream = image_stream(path, width_cm, dpi, photo)
    container_paragraph.add_run().add_picture(stream, width=Cm(width_cm), height=Cm(height_cm))


def add_figure(document, figures, path, caption, max_width, max_height, dpi=DIAGRAM_DPI, photo=False):
    width, height = fitted_size(path, max_width, max_height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(0)
    add_picture(paragraph, path, width, height, dpi, photo)
    document.add_paragraph(f"Figure {figures.next()}. {caption}", style="Figure Caption")


def add_figure_pair(document, figures, items, photo):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_borders(table, None)
    row_flags(table.rows[0])
    column = PORTRAIT_WIDTH / 2
    for index, (path, caption) in enumerate(items):
        cell = table.rows[0].cells[index]
        cell.width = Cm(column)
        width, height = fitted_size(path, column - 0.6, PORTRAIT_IMAGE_HEIGHT)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_picture(paragraph, path, width, height, PHOTO_DPI, photo)
        cell.add_paragraph(f"Figure {figures.next()}. {caption}", style="Figure Caption")
    document.add_paragraph()


def add_heading(document, text, level, page_break=None):
    paragraph = document.add_heading(text, level=level)
    if page_break is not None:
        paragraph.paragraph_format.page_break_before = page_break
    return paragraph


class Numbering:
    def __init__(self):
        self.chapter = 0
        self.section = 0

    def chapter_title(self, title):
        self.chapter += 1
        self.section = 0
        return f"{self.chapter} {title}"

    def section_title(self, title):
        self.section += 1
        return f"{self.chapter}.{self.section} {title}"


def render_blocks(document, blocks, numbering, heading_map, lead_separator=None, bullet_separator=None):
    for kind, value in blocks:
        if kind == "heading":
            level, title = value
            target = heading_map.get(level, 3)
            text = numbering.section_title(title) if target == 2 else title
            add_heading(document, text, target)
        elif kind == "paragraph":
            add_paragraph(document, reword(value))
        elif kind == "bullets":
            add_bullets(document, [reword(item) for item in value], bullet_separator)
        elif kind == "numbers":
            add_numbers(document, [(n, reword(item)) for n, item in value], lead_separator)
        elif kind == "table":
            header, *rows = value
            add_table(document, header, rows, [PORTRAIT_WIDTH / len(header)] * len(header))


def setup_header_footer(section, header_text):
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styled_run(header, header_text, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styled_run(footer, "Page ", size=9, color=MUTED)
    add_field(footer, "PAGE", "1")
    styled_run(footer, " of ", size=9, color=MUTED)
    add_field(footer, "NUMPAGES", "1")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def set_properties(document, title, subject):
    properties = document.core_properties
    properties.title = title
    properties.subject = subject
    properties.author = TEAM_NAME
    properties.last_modified_by = TEAM_NAME
    properties.comments = ""
    properties.keywords = "FreshFind, TechWiz 7, eGreen Basket, Tech Hive"
    properties.category = "Web Innovation Unleashed"


def readme_team():
    team = split_sections(read(ROOT / "README.md"), 2)
    key = next(name for name in team if name.startswith("Team"))
    return [line[2:].strip() for line in team[key].split("\n") if line.startswith("- ")]


def readme_intro():
    text = read(ROOT / "README.md")
    body = text.split("\n## ", 1)[0]
    lines = [line.strip() for line in body.split("\n") if line.strip() and not line.startswith("# ")]
    tagline = lines[0]
    return tagline, lines[1:]


def centered(document, text, size, bold=False, italic=False, color=TEXT, after=6, link=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(after)
    if link:
        add_text(paragraph, text, bold=bold, italic=italic, size=size)
    else:
        styled_run(paragraph, text, bold=bold, italic=italic, size=size, color=color)
    return paragraph


def title_page(document, team):
    logo = ROOT / "public" / "favicon.png"
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(36)
    paragraph.add_run().add_picture(str(logo), width=Cm(3.2))
    centered(document, "FreshFind", 40, bold=True, color=GREEN, after=0)
    centered(document, "Fresh All Along", 16, italic=True, color=GREEN_LIGHT, after=30)
    centered(document, "TechWiz 7 Project Report", 24, bold=True, color=TEXT, after=6)
    centered(document, f"Team {TEAM_NAME}", 16, bold=True, color=GREEN, after=24)
    centered(document, "Theme: eGreen Basket", 12, after=2)
    centered(document, "Category: Web Innovation Unleashed", 12, after=24)
    centered(document, "Team members", 12, bold=True, color=GREEN, after=4)
    for member in team:
        centered(document, member, 12, after=2)
    centered(document, "", 12, after=18)
    centered(document, "Aptech Computer Education Nigeria", 12, after=2)
    centered(document, "Advanced Diploma in Software Engineering", 12, after=2)
    centered(document, "Semester 2", 12, after=2)
    centered(document, "September 2026", 12, after=24)
    centered(document, f"Live site: {LIVE_SITE}", 12, bold=True, after=0, link=True)


def contents_page(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(18)
    styled_run(heading, "Contents", bold=True, size=20, color=GREEN)
    paragraph = document.add_paragraph()
    add_field(paragraph, 'TOC \\o "1-2" \\h \\z \\u', "Right click here and choose Update Field to build the table of contents.")


def markets_table_rows(line):
    rows = []
    for match in re.finditer(r"([A-Z][^(),:]*?) \(([^,()]+), ([^()]+)\)", line):
        rows.append([match.group(1).strip(), match.group(2).strip(), match.group(3).strip()])
    return rows


def scenario_rows(items):
    rows = []
    for number, text in items:
        parts = re.split(r"\s*Expected:\s*", text)
        if len(parts) == 2:
            did, expected = parts
        else:
            actions = [parts[0]]
            outcomes = []
            for middle in parts[1:-1]:
                outcome, _, action = middle.rpartition(". ")
                outcomes.append(outcome + ".")
                actions.append(action)
            outcomes.append(parts[-1])
            labels = [actions[0].rstrip(".").split(", ")[-1]] + [action.rstrip(".") for action in actions[1:]]
            did = actions[0].rstrip(".") + ", then " + ", then ".join(
                action[0].lower() + action[1:].rstrip(".") for action in actions[1:]
            ) + "."
            expected = " ".join(
                f"{label[0].upper() + label[1:]}: {outcome}" for label, outcome in zip(labels, outcomes)
            )
        expected = expected[0].upper() + expected[1:]
        rows.append([number, did.strip(), expected.strip()])
    return rows


def diagram_list(report_sections):
    items = []
    for kind, value in parse_blocks(report_sections["Diagrams"]):
        if kind == "numbers":
            items.extend(value)
    return items


def diagram_files():
    pictures = sorted((DOCS / "diagrams" / "png").glob("*.png"))
    explained = {path.name[:2]: path for path in (DOCS / "diagrams" / "explained").glob("*.md")}
    pairs = []
    for picture in pictures:
        prefix = picture.name[:2]
        if prefix not in explained:
            raise SystemExit(f"No explanation found for {picture.name}")
        pairs.append((picture, explained[prefix]))
    return pairs


def build_report():
    report = split_sections(read(DOCS / "PROJECT_REPORT.md"), 2)
    test_data = split_sections(read(DOCS / "TEST_DATA.md"), 2)
    lighthouse_text = read(DOCS / "LIGHTHOUSE.md")
    lighthouse = split_sections(lighthouse_text, 2)
    assumptions = read(TEAM_NOTES / "ASSUMPTIONS.md")
    roles_text = read(TEAM_NOTES / "TEAM_ROLES.md")
    ai_tools = read(DOCS / "AI_TOOLS.md")
    team = readme_team()

    document = Document()
    setup_styles(document)
    set_properties(document, "FreshFind Project Report", "TechWiz 7 Project Report by Tech Hive")
    first = document.sections[0]
    set_page(first, landscape=False)
    setup_header_footer(first, f"FreshFind Project Report  ·  Team {TEAM_NAME}")

    title_page(document, team)
    contents_page(document)

    numbering = Numbering()
    figures = Figures()
    level3_as_sections = {3: 2, 4: 3}

    add_heading(document, numbering.chapter_title("Problem Definition"), 1)
    render_blocks(document, parse_blocks(report["Problem Definition"]), numbering, level3_as_sections)

    add_heading(document, numbering.chapter_title("Design Specifications"), 1)
    design = split_sections(report["Design Specifications"], 3)
    for title, body in design.items():
        add_heading(document, numbering.section_title(title), 2)
        render_blocks(document, parse_blocks(body), numbering, level3_as_sections)
        if title == "Visual design":
            add_paragraph(document, "Figures 1 and 2 show two of the mockups, the Home page and the Market Detail page.")
            add_figure_pair(
                document,
                figures,
                [
                    (ROOT / "design" / "mockups" / "01-home.png", "Home page mockup"),
                    (ROOT / "design" / "mockups" / "03-market-detail.png", "Market Detail page mockup"),
                ],
                photo=True,
            )

    add_heading(document, numbering.chapter_title("Diagrams"), 1)
    render_blocks(document, parse_blocks(report["Diagrams"]), numbering, level3_as_sections)
    names = diagram_list(report)
    for index, (picture, explained_path) in enumerate(diagram_files()):
        name = names[index][1].split(":")[0].strip() if index < len(names) else picture.stem
        with Image.open(picture) as image:
            wide = image.width / image.height > 1.3
        if wide:
            new_section(document, landscape=True)
            add_heading(document, numbering.section_title(name), 2)
            add_figure(document, figures, picture, name, LANDSCAPE_WIDTH, LANDSCAPE_IMAGE_HEIGHT)
            new_section(document, landscape=False)
        else:
            add_heading(document, numbering.section_title(name), 2, page_break=True)
            add_figure(document, figures, picture, name, PORTRAIT_WIDTH, PORTRAIT_IMAGE_HEIGHT)
        parts = split_sections(read(explained_path), 2)
        opening = next((title for title in ("What it shows", "What a DFD is") if title in parts), None)
        for title in [opening, "The shapes", "The logic behind it"]:
            if not title or title not in parts:
                raise SystemExit(f"{explained_path.name} has no section called {title}")
            add_heading(document, title, 3)
            render_blocks(document, parse_blocks(parts[title]), numbering, {})

    add_heading(document, numbering.chapter_title("Test Data Used in the Project"), 1)
    add_heading(document, numbering.section_title("Sample data"), 2)
    for kind, value in parse_blocks(test_data["Sample data"]):
        if kind != "bullets":
            continue
        for item in value:
            rows = markets_table_rows(item)
            if len(rows) >= 2:
                lead = item.split(":")[0]
                add_paragraph(document, f"{lead[0].upper() + lead[1:]}:")
                add_table(document, ["Market", "Area", "Days and hours"], rows, [6.2, 3.8, 6.0])
            else:
                add_bullets(document, [item])
    add_heading(document, numbering.section_title("Test scenarios"), 2)
    scenarios = [block for block in parse_blocks(test_data["Test scenarios"]) if block[0] == "numbers"]
    items = [item for block in scenarios for item in block[1]]
    add_paragraph(document, "Each scenario lists what we did and the result we expected.")
    add_table(document, ["Number", "What we did", "Expected result"], scenario_rows(items), [1.8, 7.0, 7.2])

    add_heading(document, numbering.chapter_title("Testing with Google Lighthouse"), 1)
    for title, body in lighthouse.items():
        if title == "How to run it yourself":
            continue
        add_heading(document, numbering.section_title(title), 2)
        blocks = parse_blocks(body)
        images = [value for kind, value in blocks if kind == "image"]
        for kind, value in blocks:
            if kind == "paragraph":
                add_paragraph(document, reword(value))
            elif kind == "bullets":
                add_bullets(document, value, ". ")
            elif kind == "numbers":
                add_numbers(document, value, ". ")
            elif kind == "table":
                header, *rows = value
                if len(header) == 9:
                    widths = [3.6] + [1.55] * 8
                else:
                    first_width = 4.8
                    widths = [first_width] + [(PORTRAIT_WIDTH - first_width) / (len(header) - 1)] * (len(header) - 1)
                add_table(document, header, rows, widths, size=9, scores=True)
        if images:
            add_figure_pair(
                document,
                figures,
                [((DOCS / source).resolve(), alt) for alt, source in images],
                photo=True,
            )

    add_heading(document, numbering.chapter_title("Project Installation Instructions"), 1)
    render_blocks(document, parse_blocks(report["Project Installation Instructions"]), numbering, level3_as_sections)

    add_heading(document, numbering.chapter_title("Assumptions"), 1)
    add_paragraph(document, "We made these assumptions while building FreshFind.")
    render_blocks(document, parse_blocks(intro_of(assumptions)), numbering, {})

    add_heading(document, numbering.chapter_title("Team and Challenges"), 1)
    add_paragraph(
        document,
        f"FreshFind was built by Team {TEAM_NAME}, four students of Aptech Computer Education Nigeria. "
        "Each member was responsible for their own pages, and the team leader was also responsible for the parts every page shares.",
    )
    roles = split_sections(split_sections(roles_text, 2)["Roles"], 3)
    add_heading(document, numbering.section_title("Roles"), 2)
    for person in roles:
        if person not in ROLE_TEXT:
            raise SystemExit(f"No role text for {person}")
        add_heading(document, person, 3)
        add_paragraph(document, ROLE_TEXT[person])
    challenges_title = next(title for title in split_sections(roles_text, 2) if title.startswith("Challenges"))
    add_heading(document, numbering.section_title(challenges_title), 2)
    challenge_blocks = parse_blocks(split_sections(roles_text, 2)[challenges_title])
    for kind, value in challenge_blocks:
        if kind == "numbers":
            add_numbers(document, [(n, reword(item)) for n, item in value], ". ")
        elif kind == "paragraph":
            add_paragraph(document, reword(value))

    add_heading(document, numbering.chapter_title("AI Tools Acknowledgment"), 1)
    render_blocks(document, parse_blocks(ai_tools.split("\n", 1)[1]), numbering, {2: 3}, bullet_separator=": ")

    document.save(REPORT_PATH)
    return document


def build_readme():
    tagline, paragraphs = readme_intro()
    readme = split_sections(read(ROOT / "README.md"), 2)
    assumptions = read(TEAM_NOTES / "ASSUMPTIONS.md")
    team = readme_team()

    document = Document()
    setup_styles(document)
    document.styles["Heading 1"].paragraph_format.page_break_before = False
    document.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    set_properties(document, "FreshFind ReadMe", "ReadMe for the FreshFind source code by Tech Hive")
    section = document.sections[0]
    set_page(section, landscape=False)
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    styled_run(section.footer.paragraphs[0], "FreshFind ReadMe  ·  Page ", size=9, color=MUTED)
    add_field(section.footer.paragraphs[0], "PAGE", "1")

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(0)
    styled_run(title, "FreshFind ReadMe", bold=True, size=28, color=GREEN)
    byline = document.add_paragraph()
    byline.paragraph_format.space_after = Pt(4)
    styled_run(byline, f"By Team {TEAM_NAME}", bold=True, size=13, color=GREEN_LIGHT)
    motto = document.add_paragraph()
    motto.paragraph_format.space_after = Pt(14)
    styled_run(motto, tagline, italic=True, size=12, color=MUTED)

    add_heading(document, "About the project", 1)
    for text in paragraphs:
        add_paragraph(document, text)

    add_heading(document, "Live site", 1)
    live = readme["Live site"].strip()
    add_paragraph(document, live)

    add_heading(document, "Run it on your computer", 1)
    for kind, value in parse_blocks(readme["Run it on your computer"]):
        if kind == "paragraph":
            add_paragraph(document, value)
        elif kind == "numbers":
            add_numbers(document, value)
        elif kind == "bullets":
            add_bullets(document, value)

    add_heading(document, "Assumptions", 1)
    render_blocks(document, parse_blocks(intro_of(assumptions)), Numbering(), {})

    add_heading(document, f"Team {TEAM_NAME}", 1)
    add_bullets(document, team)
    school = readme.get("School")
    if school:
        add_heading(document, "School", 1)
        for line in school.split("\n"):
            if line.strip():
                add_paragraph(document, line.strip())

    document.save(README_DOCX)
    return document


def document_text(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def check_text(label, text):
    problems = []
    for name in PRIVATE_NAMES:
        if name in text:
            problems.append(f"mentions {name}")
    for pattern in CODE_PATTERNS:
        match = re.search(pattern, text)
        if match:
            problems.append(f"looks like code: {match.group(0)!r}")
    for problem in problems:
        print(f"WARNING {label}: {problem}")
    return not problems


WORD_SCRIPT = r"""
param([string]$Report, [string]$ReadMeDocx, [string]$ReadMeDoc)
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
try {
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $doc = $word.Documents.Open($Report)
  foreach ($id in -20, -21) {
    $style = $doc.Styles.Item($id)
    $style.Font.Size = 10
    $style.ParagraphFormat.LineSpacingRule = 0
    $style.ParagraphFormat.SpaceAfter = 1
    if ($id -eq -20) { $style.Font.Bold = $true; $style.ParagraphFormat.SpaceBefore = 3 }
  }
  foreach ($toc in $doc.TablesOfContents) { $toc.Update() }
  $doc.Fields.Update() | Out-Null
  foreach ($toc in $doc.TablesOfContents) { $toc.UpdatePageNumbers() }
  $doc.Save()
  $doc.Close()
  $doc = $word.Documents.Open($ReadMeDocx)
  $target = $ReadMeDoc
  $format = 0
  $doc.SaveAs([ref]$target, [ref]$format)
  $doc.Close()
}
finally {
  $word.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}
"""


def run_word():
    if sys.platform != "win32" or not shutil.which("powershell"):
        return False
    with tempfile.TemporaryDirectory() as folder:
        script = Path(folder) / "word.ps1"
        script.write_text(WORD_SCRIPT, encoding="utf-8")
        result = subprocess.run(
            [
                "powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                "-File", str(script),
                "-Report", str(REPORT_PATH),
                "-ReadMeDocx", str(README_DOCX),
                "-ReadMeDoc", str(README_DOC),
            ],
            capture_output=True,
            text=True,
        )
    if result.returncode != 0:
        print("Word automation failed:", (result.stderr or result.stdout).strip()[:500])
        return False
    return README_DOC.exists()


def find_libreoffice():
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    ):
        if candidate.exists():
            return str(candidate)
    return None


def run_libreoffice():
    soffice = find_libreoffice()
    if not soffice:
        return False
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "doc", "--outdir", str(ROOT), str(README_DOCX)],
        capture_output=True,
        text=True,
    )
    return result.returncode == 0 and README_DOC.exists()


def ask_word_to_update_fields():
    document = Document(REPORT_PATH)
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.save(REPORT_PATH)


def main():
    report = build_report()
    readme = build_readme()
    report_ok = check_text("report", document_text(report))
    readme_ok = check_text("ReadMe", document_text(readme))

    if README_DOC.exists():
        README_DOC.unlink()
    if run_word():
        README_DOCX.unlink()
        print("Word filled in the table of contents and saved ReadMe.doc in the Word 97-2003 format.")
    else:
        ask_word_to_update_fields()
        if run_libreoffice():
            README_DOCX.unlink()
            print("LibreOffice saved ReadMe.doc. Word will offer to fill in the table of contents when the report opens.")
        else:
            print("Neither Word nor LibreOffice is available, so ReadMe.docx was kept.")

    print(f"Report: {REPORT_PATH.relative_to(ROOT)}")
    print(f"ReadMe: {(README_DOC if README_DOC.exists() else README_DOCX).relative_to(ROOT)}")
    if not (report_ok and readme_ok):
        print("Check the warnings above.")


if __name__ == "__main__":
    main()
